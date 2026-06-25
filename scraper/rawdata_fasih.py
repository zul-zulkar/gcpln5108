"""
rawdata_fasih.py — Scraper raw data tabel FASIH BPS

Prosedur:
  1. Filter status: aktifkan SUBMITTED BY PENCACAH saja
  2. Rows per page: 100
  3. Iterasi UPI → UP3 → ULP (Seririt / Singaraja / Tejakula) → Pencacah
  4. Ekstrak semua baris tabel + paginasi
  5. Simpan ke rawdata/ sebagai CSV (streaming) lalu konversi ke Excel

Checkpoint: rawdata/checkpoint.json — bisa resume jika terputus.
"""

import asyncio
import csv
import json
import os
from datetime import datetime
from playwright.async_api import async_playwright
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

# ─── Config ───────────────────────────────────────────────────────────────────
RAWDATA_DIR     = "rawdata"
CHECKPOINT_FILE = os.path.join(RAWDATA_DIR, "checkpoint.json")

SURVEY_LIST_URL = "https://fasih-sm.bps.go.id/survey-collection/survey"
SURVEYS = {
    "pascabayar": "https://fasih-sm.bps.go.id/survey-collection/collect/2e31188c-a617-4163-8056-edccf93d8d79",
    "prabayar":   "https://fasih-sm.bps.go.id/survey-collection/collect/2395b67d-d1af-4739-9ef8-c0cc0aa9ce9a",
}

UPI_TEXT      = "[55]"
UP3_TEXT      = "[55UTR]"
ULP_WHITELIST = ["seririt", "singaraja", "tejakula"]   # case-insensitive contains
ROWS_PER_PAGE = 100

# Extra kolom yang ditambahkan di depan setiap baris tabel
EXTRA_HEADERS = ["Survey", "ULP", "Pencacah"]

# ─── Login ────────────────────────────────────────────────────────────────────

async def login(page, username, password):
    print(f"[LOGIN] {page.url}")
    try:
        btn = await page.wait_for_selector("a.login-button", timeout=10000)
        await btn.click()
        await page.wait_for_load_state("domcontentloaded", timeout=20000)
        print(f"[LOGIN] SSO: {page.url}")
    except Exception as e:
        print(f"[LOGIN] SSO btn error: {e}")
    try:
        await page.wait_for_selector("input[name='username']", timeout=10000)
        await page.fill("input[name='username']", username)
        await page.fill("input[name='password']", password)
        await page.click("input[type='submit']")
        await page.wait_for_load_state("domcontentloaded", timeout=20000)
        print(f"[LOGIN] Done: {page.url}")
    except Exception as e:
        print(f"[LOGIN] Form error: {e}")


# ─── ngx-select helpers ───────────────────────────────────────────────────────

async def ngx_select(page, selector, search_text, timeout=10000):
    """Pilih opsi di ngx-select. Return teks yang ter-select atau None."""
    toggle_sel      = f"{selector} .ngx-select__toggle"
    search_input_sel = f"{selector} input.ngx-select__search"

    toggle = await page.wait_for_selector(toggle_sel, timeout=timeout)
    await page.evaluate(
        "el => el.dispatchEvent(new MouseEvent('click', {bubbles:true,cancelable:true}))",
        toggle,
    )
    await page.wait_for_timeout(400)

    search_input = await page.query_selector(search_input_sel)
    if search_input:
        await page.evaluate(
            """([el]) => {
                el.focus(); el.value = '';
                el.dispatchEvent(new Event('input', {bubbles:true}));
            }""",
            [search_input],
        )
        await page.wait_for_timeout(200)
        await page.evaluate(
            """([el, val]) => {
                el.value = val;
                el.dispatchEvent(new Event('input', {bubbles:true}));
                el.dispatchEvent(new KeyboardEvent('keyup', {bubbles:true}));
            }""",
            [search_input, search_text],
        )
        await page.wait_for_timeout(800)

    clicked = await page.evaluate(
        """([sel, search]) => {
            const items = document.querySelectorAll(sel + ' .ngx-select__item');
            const q = search.toLowerCase();
            for (const item of items) {
                const t = (item.innerText || '').trim();
                if (t.toLowerCase().includes(q)) { item.click(); return t; }
            }
            return null;
        }""",
        [selector, search_text],
    )

    if clicked:
        print(f"    Selected: {clicked}")
        return clicked
    print(f"    [WARN] Tidak ada match '{search_text}' di {selector}")
    await page.keyboard.press("Escape")
    return None


async def ngx_get_options(page, selector, timeout=8000):
    """Buka ngx-select dan kembalikan semua teks opsi yang tersedia."""
    toggle_sel = f"{selector} .ngx-select__toggle"
    try:
        toggle = await page.wait_for_selector(toggle_sel, timeout=timeout)
        await page.evaluate(
            "el => el.dispatchEvent(new MouseEvent('click', {bubbles:true,cancelable:true}))",
            toggle,
        )
        await page.wait_for_timeout(600)
        options = await page.evaluate(
            f"""() => {{
                const items = document.querySelectorAll('{selector} .ngx-select__item');
                return Array.from(items)
                    .map(el => (el.innerText || '').trim())
                    .filter(t => t);
            }}"""
        )
        await page.keyboard.press("Escape")
        await page.wait_for_timeout(200)
        return options
    except Exception as e:
        print(f"    [WARN] ngx_get_options({selector}): {e}")
        try:
            await page.keyboard.press("Escape")
        except Exception:
            pass
        return []


# ─── Filter sidebar ───────────────────────────────────────────────────────────

async def open_filter(page):
    try:
        btn = await page.wait_for_selector("button:has-text('Filter')", timeout=10000)
        await page.evaluate("el => el.click()", btn)
        await page.wait_for_selector(".sidebar-content", timeout=5000)
        return True
    except Exception:
        return False


async def click_filter_data(page):
    try:
        btn = await page.wait_for_selector(
            ".sidebar-content button:has-text('Filter Data')", timeout=8000
        )
        await page.evaluate("el => el.click()", btn)
        try:
            await page.wait_for_load_state("networkidle", timeout=6000)
        except Exception:
            await page.wait_for_timeout(1000)
    except Exception as e:
        print(f"    [WARN] click_filter_data: {e}")


async def ensure_sidebar_open(page):
    """Pastikan filter sidebar terbuka; buka kembali jika perlu."""
    visible = await page.evaluate(
        """() => {
            const s = document.querySelector('.sidebar-content');
            return s ? (s.offsetParent !== null || s.getBoundingClientRect().width > 0) : false;
        }"""
    )
    if not visible:
        return await open_filter(page)
    return True


# ─── Status badge filter ──────────────────────────────────────────────────────

async def activate_submitted_filter(page):
    """Klik badge SUBMITTED BY PENCACAH jika belum aktif."""
    try:
        btn = await page.wait_for_selector(
            "button.btn-outline-primary:has-text('SUBMITTED')", timeout=8000
        )
        # Cek apakah sudah aktif (class active / btn-primary)
        already = await page.evaluate(
            "el => el.classList.contains('active') || el.classList.contains('btn-primary')",
            btn,
        )
        if not already:
            await page.evaluate("el => el.click()", btn)
            try:
                await page.wait_for_load_state("networkidle", timeout=6000)
            except Exception:
                await page.wait_for_timeout(1000)
            print("    [STATUS] Filter SUBMITTED BY PENCACAH diaktifkan")
        else:
            print("    [STATUS] Filter SUBMITTED sudah aktif")
    except Exception as e:
        print(f"    [WARN] activate_submitted_filter: {e}")


# ─── Rows per page ────────────────────────────────────────────────────────────

async def set_rows_per_page(page, n=100):
    """Set jumlah baris per halaman. Return True jika berhasil."""
    result = await page.evaluate(
        """(n) => {
            // Cari <select> yang berisi opsi halaman (10/25/50/100)
            for (const sel of document.querySelectorAll('select')) {
                const vals = Array.from(sel.options).map(o => parseInt(o.value)).filter(v => !isNaN(v));
                if (vals.some(v => v >= 10 && v <= 500) && vals.length >= 2) {
                    // Cari nilai terdekat
                    let best = null, minDiff = Infinity;
                    for (const opt of sel.options) {
                        const v = parseInt(opt.value);
                        if (!isNaN(v) && Math.abs(v - n) < minDiff) {
                            minDiff = Math.abs(v - n);
                            best = opt;
                        }
                    }
                    if (best) {
                        sel.value = best.value;
                        sel.dispatchEvent(new Event('change', {bubbles:true}));
                        return 'select:' + best.value;
                    }
                }
            }
            return null;
        }""",
        n,
    )
    if result:
        print(f"    [ROWS/PAGE] {result}")
        try:
            await page.wait_for_load_state("networkidle", timeout=5000)
        except Exception:
            await page.wait_for_timeout(800)
        return True
    print(f"    [WARN] Tidak bisa set rows per page ke {n}")
    return False


# ─── Tabel & Paginasi ─────────────────────────────────────────────────────────

async def extract_table(page):
    """Ekstrak header dan baris dari tabel yang terlihat di halaman."""
    return await page.evaluate(
        """() => {
            const table = document.querySelector('table');
            if (!table) return {headers: [], rows: []};
            const headers = Array.from(
                table.querySelectorAll('thead tr:last-child th')
            ).map(th => (th.innerText || th.textContent || '').replace(/\\n/g,' ').trim());
            const rows = Array.from(table.querySelectorAll('tbody tr'))
                .map(tr =>
                    Array.from(tr.querySelectorAll('td'))
                        .map(td => (td.innerText || td.textContent || '').replace(/\\n/g,' ').trim())
                )
                .filter(r => r.length > 0 && r.some(c => c));
            return {headers, rows};
        }"""
    )


async def has_next_page(page):
    """Cek apakah ada tombol next page yang aktif."""
    return await page.evaluate(
        """() => {
            // Angular Material paginator
            const mat = document.querySelector(
                'button.mat-paginator-navigation-next, button[aria-label="Next page"]'
            );
            if (mat) return !mat.disabled;

            // Bootstrap pagination
            const bs = document.querySelector(
                'li.page-item.next:not(.disabled) a, li.pagination-next:not(.disabled) a'
            );
            if (bs) return true;

            // Generic: cari tombol › atau »
            for (const el of document.querySelectorAll('button, a.page-link')) {
                const t = (el.innerText || el.textContent || '').trim();
                const lbl = (el.getAttribute('aria-label') || '').toLowerCase();
                if ((t === '›' || t === '»' || t === '>' || lbl.includes('next'))
                        && !el.disabled && !el.closest('.disabled')) {
                    return true;
                }
            }
            return false;
        }"""
    )


async def click_next_page(page):
    """Klik tombol next page. Return True jika berhasil."""
    clicked = await page.evaluate(
        """() => {
            const mat = document.querySelector(
                'button.mat-paginator-navigation-next, button[aria-label="Next page"]'
            );
            if (mat && !mat.disabled) { mat.click(); return 'mat'; }

            const bs = document.querySelector(
                'li.page-item.next:not(.disabled) a, li.pagination-next:not(.disabled) a'
            );
            if (bs) { bs.click(); return 'bs'; }

            for (const el of document.querySelectorAll('button, a.page-link')) {
                const t = (el.innerText || el.textContent || '').trim();
                const lbl = (el.getAttribute('aria-label') || '').toLowerCase();
                if ((t === '›' || t === '»' || t === '>' || lbl.includes('next'))
                        && !el.disabled && !el.closest('.disabled')) {
                    el.click(); return 'generic';
                }
            }
            return null;
        }"""
    )
    if clicked:
        try:
            await page.wait_for_load_state("networkidle", timeout=8000)
        except Exception:
            await page.wait_for_timeout(1200)
        return True
    return False


# ─── Checkpoint ───────────────────────────────────────────────────────────────

def load_checkpoint():
    if os.path.exists(CHECKPOINT_FILE):
        try:
            with open(CHECKPOINT_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def save_checkpoint(cp):
    os.makedirs(RAWDATA_DIR, exist_ok=True)
    with open(CHECKPOINT_FILE, "w", encoding="utf-8") as f:
        json.dump(cp, f, indent=2, ensure_ascii=False)


# ─── CSV helpers ──────────────────────────────────────────────────────────────

def write_csv_header(csv_path, headers):
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerow(headers)


def append_rows_csv(csv_path, rows, prefix):
    """Append rows ke CSV, tambahkan kolom prefix di depan setiap baris."""
    with open(csv_path, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        for row in rows:
            w.writerow(prefix + list(row))


# ─── Konversi CSV → Excel (write-only, hemat memori) ─────────────────────────

def build_excel_from_csvs(csv_map, xlsx_path):
    """
    csv_map: {sheet_name: csv_path}
    Gunakan write-only workbook agar tidak load semua baris ke RAM.
    """
    print(f"\n[CONVERT] Membangun Excel: {xlsx_path}")
    hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)

    wb = openpyxl.Workbook(write_only=True)
    for sheet_name, csv_path in csv_map.items():
        if not os.path.exists(csv_path):
            print(f"  [WARN] Tidak ada CSV: {csv_path}")
            continue
        ws = wb.create_sheet(title=sheet_name)
        with open(csv_path, "r", encoding="utf-8-sig") as f:
            for i, row in enumerate(csv.reader(f)):
                if i == 0:
                    cells = []
                    for val in row:
                        c = WriteOnlyCell(ws, value=val)
                        c.fill = hdr_fill
                        c.font = hdr_font
                        cells.append(c)
                    ws.append(cells)
                else:
                    ws.append(row)
        print(f"  [SHEET] {sheet_name} selesai")
    wb.save(xlsx_path)
    print(f"[SAVED] {xlsx_path}")
    return xlsx_path


# ─── Core scraper per survey ──────────────────────────────────────────────────

async def scrape_one_survey(page, survey_name, url, username, password,
                             csv_path, checkpoint, stop_event=None):
    print(f"\n{'='*70}")
    print(f"[SURVEY] {survey_name.upper()}")

    # ── Navigasi ──────────────────────────────────────────────────────────────
    collect_id = url.split("/")[-1]
    if collect_id not in page.url:
        print(f"[NAV] → {url}")
        await page.goto(url, wait_until="domcontentloaded", timeout=90000)
        if "oauth" in page.url or "sso" in page.url:
            print("[LOGIN] Session expired, re-login…")
            await login(page, username, password)
            await page.goto(url, wait_until="domcontentloaded", timeout=90000)
        try:
            await page.wait_for_load_state("networkidle", timeout=15000)
        except Exception:
            await page.wait_for_timeout(2000)

    # ── Tab Data ──────────────────────────────────────────────────────────────
    try:
        tab = await page.wait_for_selector("li.nav-item a:has-text('Data')", timeout=10000)
        await page.evaluate("el => el.click()", tab)
        await page.wait_for_timeout(500)
    except Exception:
        pass

    await page.wait_for_selector("button:has-text('Filter')", timeout=25000)

    # ── Init filter: UPI + UP3 ────────────────────────────────────────────────
    if not await open_filter(page):
        print("[ERROR] Tidak bisa buka filter sidebar")
        return

    await ngx_select(page, 'ngx-select[name="region1Id"]', UPI_TEXT)
    await ngx_select(page, 'ngx-select[name="region2Id"]', UP3_TEXT)
    await click_filter_data(page)
    try:
        await page.wait_for_load_state("networkidle", timeout=6000)
    except Exception:
        await page.wait_for_timeout(1500)

    # ── Enumerate ULP ─────────────────────────────────────────────────────────
    print("[ULP] Mengambil daftar ULP…")
    ulp_all = await ngx_get_options(page, 'ngx-select[name="region3Id"]')
    ulp_list = [u for u in ulp_all if any(w in u.lower() for w in ULP_WHITELIST)]
    print(f"[ULP] Total: {len(ulp_all)} | Whitelist match: {ulp_list}")

    cp_survey   = checkpoint.setdefault(survey_name, {})
    hdr_written = bool(cp_survey.get("_headers_written"))
    total_rows  = cp_survey.get("_total_rows", 0)

    for ulp_text in ulp_list:
        if stop_event and stop_event.is_set():
            print("[STOP] Dihentikan.")
            return

        ulp_key  = ulp_text[:40]
        cp_ulp   = cp_survey.setdefault(ulp_key, {})

        print(f"\n  [ULP] ▶ {ulp_text}")

        # Pastikan sidebar buka, pilih ULP
        await ensure_sidebar_open(page)
        sel = await ngx_select(page, 'ngx-select[name="region3Id"]', ulp_text)
        if not sel:
            print(f"  [SKIP-ULP] Tidak bisa pilih ULP: {ulp_text}")
            continue
        await click_filter_data(page)
        try:
            await page.wait_for_load_state("networkidle", timeout=6000)
        except Exception:
            await page.wait_for_timeout(1500)

        # Enumerate Pencacah untuk ULP ini
        print(f"  [PENCACAH] Mengambil daftar pencacah…")
        penc_list = await ngx_get_options(page, 'ngx-select[optionvaluefield="username"]')
        print(f"  [PENCACAH] {len(penc_list)} pencacah ditemukan")

        for penc_text in penc_list:
            if stop_event and stop_event.is_set():
                print("[STOP] Dihentikan.")
                return

            penc_key = penc_text[:60]
            if cp_ulp.get(penc_key) == "done":
                print(f"    [SKIP] {penc_text} (sudah selesai)")
                continue

            print(f"    [PENCACAH] ▶ {penc_text}")

            try:
                # Pastikan sidebar masih buka
                await ensure_sidebar_open(page)

                # Pilih pencacah
                sel_p = await ngx_select(
                    page, 'ngx-select[optionvaluefield="username"]', penc_text
                )
                if not sel_p:
                    print(f"    [WARN] Pencacah tidak ter-select: {penc_text}")
                    cp_ulp[penc_key] = "skip:no_match"
                    save_checkpoint(checkpoint)
                    continue

                await click_filter_data(page)

                # Aktifkan filter SUBMITTED BY PENCACAH
                await activate_submitted_filter(page)

                # Set 100 rows/page (panggil sekali; persist selama session)
                if not cp_survey.get("_rows_per_page_set"):
                    ok = await set_rows_per_page(page, ROWS_PER_PAGE)
                    if ok:
                        cp_survey["_rows_per_page_set"] = True
                        save_checkpoint(checkpoint)

                # Tunggu tabel muncul
                try:
                    await page.wait_for_selector("table tbody tr", timeout=8000)
                except Exception:
                    print(f"    [INFO] Tidak ada baris untuk {penc_text}")
                    cp_ulp[penc_key] = "done"
                    save_checkpoint(checkpoint)
                    continue

                # Tulis header CSV (sekali saja per survey)
                if not hdr_written:
                    tbl0 = await extract_table(page)
                    if tbl0["headers"]:
                        write_csv_header(
                            csv_path, EXTRA_HEADERS + tbl0["headers"]
                        )
                        hdr_written = True
                        cp_survey["_headers_written"] = True
                        save_checkpoint(checkpoint)

                # Paginate dan ekstrak
                page_num   = 0
                penc_rows  = 0
                while True:
                    page_num += 1
                    tbl = await extract_table(page)
                    rows = tbl["rows"]
                    if not rows:
                        break

                    prefix = [survey_name, ulp_text, penc_text]
                    append_rows_csv(csv_path, rows, prefix)

                    penc_rows  += len(rows)
                    total_rows += len(rows)
                    print(
                        f"      Hal {page_num}: {len(rows)} baris "
                        f"(pencacah={penc_rows}, total={total_rows})"
                    )

                    if not await has_next_page(page):
                        break
                    await click_next_page(page)

                print(f"    ✓ {penc_text}: {penc_rows} baris")
                cp_ulp[penc_key] = "done"
                cp_survey["_total_rows"] = total_rows
                save_checkpoint(checkpoint)

            except Exception as e:
                print(f"    [ERROR] {penc_text}: {e}")
                cp_ulp[penc_key] = f"error:{str(e)[:80]}"
                save_checkpoint(checkpoint)
                # Recover: buka ulang filter dan re-set ULP
                try:
                    await page.keyboard.press("Escape")
                    await page.wait_for_timeout(500)
                    await ensure_sidebar_open(page)
                    await ngx_select(page, 'ngx-select[name="region3Id"]', ulp_text)
                except Exception:
                    pass

    cp_survey["_total_rows"] = total_rows
    save_checkpoint(checkpoint)
    print(f"\n[DONE] {survey_name}: {total_rows} baris → {csv_path}")


# ─── Entry point ──────────────────────────────────────────────────────────────

async def main_rawdata(username, password, stop_event=None, headless=False):
    os.makedirs(RAWDATA_DIR, exist_ok=True)
    timestamp  = datetime.now().strftime("%Y%m%d_%H%M%S")
    checkpoint = load_checkpoint()

    # Tentukan path CSV (resume jika sudah ada di checkpoint)
    csv_paths = {}
    for name in SURVEYS:
        key = f"_csv_{name}"
        prev = checkpoint.get(key)
        if prev and os.path.exists(prev):
            csv_paths[name] = prev
            print(f"[RESUME] {name}: melanjutkan dari {prev}")
        else:
            csv_paths[name] = os.path.join(RAWDATA_DIR, f"rawdata_{name}_{timestamp}.csv")
            checkpoint[key] = csv_paths[name]
            save_checkpoint(checkpoint)

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=headless, slow_mo=30)
        context = await browser.new_context(viewport={"width": 1400, "height": 900})
        page    = await context.new_page()

        print(f"[INFO] Opening: {SURVEY_LIST_URL}")
        await page.goto(SURVEY_LIST_URL, wait_until="domcontentloaded", timeout=30000)

        if SURVEY_LIST_URL not in page.url:
            await login(page, username, password)
            try:
                await page.wait_for_load_state("networkidle", timeout=20000)
            except Exception:
                await page.wait_for_timeout(4000)
        print(f"[INFO] Logged in: {page.url}")

        for survey_name, url in SURVEYS.items():
            if stop_event and stop_event.is_set():
                break
            await scrape_one_survey(
                page, survey_name, url, username, password,
                csv_paths[survey_name], checkpoint, stop_event,
            )

        await browser.close()
        print("[INFO] Browser ditutup.")

    # Konversi CSV → Excel (write-only, hemat memori)
    if not (stop_event and stop_event.is_set()):
        xlsx_path = os.path.join(RAWDATA_DIR, f"rawdata_fasih_{timestamp}.xlsx")
        sheet_map = {
            ("Pascabayar" if k == "pascabayar" else "Prabayar"): v
            for k, v in csv_paths.items()
        }
        build_excel_from_csvs(sheet_map, xlsx_path)
        checkpoint["_xlsx"] = xlsx_path
        save_checkpoint(checkpoint)
        return xlsx_path
    return None


if __name__ == "__main__":
    u = "m.zulkarnain"
    pw = "0n3g41Z*mbra"
    asyncio.run(main_rawdata(u, pw))
